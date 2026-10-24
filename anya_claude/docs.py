# -*- coding: utf-8 -*-
"""文件面板（階段 3）：使用者上傳的文件整份放進對話，回答附頁碼引用。純函式，無 streamlit 依賴。

做法（計畫檔階段 3；spike T5：10 條引文用 pypdf 核對，10/10 在標示的頁面上）：
- PDF 原檔 base64 送：Claude 每頁同時讀文字和影像，掃描檔也看得懂（但掃描檔沒有文字層，引用不了）。
  引用帶頁碼（page_location，end_page_number 不含）。
- 純文字：原樣送。引用帶字元位置，沒有頁碼。
- Word、簡報、Excel、CSV：每段、每張投影片、每一段列做成一個區塊（custom content），引用帶區塊編號，
  這裡換成「第 2 段」「第 3 張投影片」「工作表「營收」第 2–41 列」。
- 全部放在對話最前面的一則訊息（文件架），最後一份標 1h 快取斷點：之後每題、每次模型呼叫都讀快取。
- 不用 FAISS／embedding、不做 OCR（舊版 docstore 的做法）。簡報用標準函式庫解析，不加 python-pptx。

官方限制（2026-10 文件）：一次請求最多 600 頁 PDF、整包 32MB（base64 會多三成）。
"""
from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from langchain_core.messages import AIMessage, BaseMessage, HumanMessage

MAX_DOCS = 5
MAX_PDF_PAGES = 300                 # 所有 PDF 合計；Haiku 的 prompt 超過 10 萬 token 就改用約 5 倍的價目
MAX_TOTAL_BYTES = 20 * 1024 * 1024  # base64 後約 27MB，留空間給圖片與對話
MAX_TEXT_CHARS = 200_000            # 每份文字類文件
ROWS_PER_BLOCK = 40                 # Excel／CSV 每個引用區塊的列數（含表頭另外重複一次）
SCAN_CHECK_PAGES = 3                # 前幾頁都抽不到字 → 當作掃描檔
CACHE_DOCS = {"type": "ephemeral", "ttl": "1h"}
SHELF_KEY = "anya_doc_shelf"        # 文件架訊息的標記（評分、上一題摘要要跳過它）

TEXT_EXTS = {".txt", ".md"}
OLD_EXTS = {".doc", ".xls", ".ppt"}
SUPPORTED = (".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md")


@dataclass
class Doc:
    name: str
    kind: str                       # pdf | text | blocks
    size: int                       # 原檔位元組
    pages: int = 0                  # PDF 頁數
    data: str = ""                  # pdf：base64；text：全文
    blocks: list[str] = field(default_factory=list)   # blocks：每個引用區塊的文字
    labels: list[str] = field(default_factory=list)   # blocks：每個區塊的位置名稱
    scanned: bool = False           # PDF 抽不到文字層：看得懂，但引用不了
    amount: str = ""                # 面板上顯示的份量（「3 張投影片」「85 列資料」「12 段」）

    @property
    def summary(self) -> str:
        if self.kind == "pdf":
            return f"{self.name}（{self.pages} 頁{'，掃描檔' if self.scanned else ''}）"
        if self.kind == "text":
            return f"{self.name}（{len(self.data):,} 字）"
        return f"{self.name}（{self.amount or f'{len(self.blocks)} 段'}）"


def _ext(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m.group(0).lower() if m else ""


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp950", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(name: str, data: bytes) -> tuple[Doc | None, str]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            return None, f"「{name}」有密碼保護，請先解除密碼再上傳。"
        pages = len(reader.pages)
        head = "".join((reader.pages[i].extract_text() or "") for i in range(min(pages, SCAN_CHECK_PAGES)))
    except Exception as e:  # noqa: BLE001 — 壞檔、不是真的 PDF
        return None, f"「{name}」打不開（{type(e).__name__}），請確認是正常的 PDF。"
    doc = Doc(name, "pdf", len(data), pages=pages, data=base64.standard_b64encode(data).decode(),
              scanned=not head.strip())
    return doc, ("這份看起來是掃描檔：安妮亞讀得懂，但回答沒辦法附頁碼引用。" if doc.scanned else "")


def _docx(data: bytes) -> tuple[list[str], list[str], int]:
    """每段一個區塊、每個表格一個區塊（表格排在段落後面）。2026-10-24：整份當純文字送時，引用常把整份文字
    標成一條，看不出引的是哪一條規定；一段一塊，引用就是「第 2 段」那一段。"""
    import docx
    d = docx.Document(io.BytesIO(data))
    blocks, labels = [], []
    for p in d.paragraphs:
        if p.text.strip():
            labels.append(f"第 {len(labels) + 1} 段")
            blocks.append(p.text.strip())
    n_paras = len(blocks)
    for k, t in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        rows = [r for r in rows if r.strip(" |")]
        if rows:
            labels.append(f"表格 {k}")
            blocks.append("\n".join(rows))
    return blocks, labels, n_paras


_SLIDE_RE = re.compile(r"ppt/slides/slide(\d+)\.xml$")
_PARA_RE = re.compile(r"<a:p[ >].*?</a:p>", re.S)
_RUN_RE = re.compile(r"<a:t>(.*?)</a:t>", re.S)


def _unescape(s: str) -> str:
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
            .replace("&apos;", "'").replace("&amp;", "&"))


def _pptx(data: bytes) -> tuple[list[str], list[str]]:
    """每張投影片一個區塊（照投影片編號排序）。只取文字框的字，備忘稿不取。"""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        slides = sorted(((int(m.group(1)), n) for n in z.namelist() if (m := _SLIDE_RE.match(n))))
        blocks, labels = [], []
        for no, path in slides:
            xml = z.read(path).decode("utf-8", errors="replace")
            lines = ["".join(_unescape(t) for t in _RUN_RE.findall(p)).strip() for p in _PARA_RE.findall(xml)]
            body = "\n".join(ln for ln in lines if ln) or "（這張投影片沒有文字）"
            labels.append(f"第 {no} 張投影片")
            blocks.append(f"【第 {no} 張投影片】\n{body}")
    return blocks, labels


def _row_blocks(rows: list[list[str]], where: str) -> tuple[list[str], list[str]]:
    """表格每 ROWS_PER_BLOCK 列一個區塊，每塊都重複表頭（第 1 列），引用時看得出欄位。"""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    head, body = rows[0], rows[1:] or []
    blocks, labels = [], []
    if not body:
        return [f"【{where}第 1 列】\n" + "\t".join(head)], [f"{where}第 1 列"]
    for i in range(0, len(body), ROWS_PER_BLOCK):
        chunk = body[i:i + ROWS_PER_BLOCK]
        first, last = i + 2, i + 1 + len(chunk)          # 試算表的列號：表頭是第 1 列
        label = f"{where}第 {first}–{last} 列" if last > first else f"{where}第 {first} 列"
        labels.append(label)
        blocks.append(f"【{label}】\n" + "\n".join("\t".join(r) for r in [head, *chunk]))
    return blocks, labels


def _xlsx(data: bytes) -> tuple[list[str], list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks, labels = [], []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        b, lb = _row_blocks(rows, f"工作表「{ws.title}」")
        blocks += b
        labels += lb
    wb.close()
    return blocks, labels


def _csv(data: bytes) -> tuple[list[str], list[str]]:
    return _row_blocks(list(csv.reader(io.StringIO(_decode(data)))), "")


def _cap_blocks(blocks: list[str], labels: list[str]) -> tuple[list[str], list[str], bool]:
    total, keep = 0, 0
    for b in blocks:
        if total + len(b) > MAX_TEXT_CHARS:
            break
        total += len(b)
        keep += 1
    if keep == 0:                                   # 第一塊就超過：截斷它，至少留一塊
        return [blocks[0][:MAX_TEXT_CHARS]], labels[:1], True
    return blocks[:keep], labels[:keep], keep < len(blocks)


def prepare(name: str, data: bytes, existing: list[Doc]) -> tuple[Doc | None, str]:
    """把上傳的檔案整理成 Doc。回傳 (Doc 或 None, 給使用者看的說明)：None 時說明是錯誤，否則是提醒（可能空）。"""
    ext = _ext(name)
    if ext in OLD_EXTS:
        return None, f"「{name}」是舊格式，請另存成 {ext}x 或 PDF 再上傳。"
    if ext not in SUPPORTED:
        return None, f"「{name}」的格式不支援（可以上傳 PDF、Word、Excel、PowerPoint、CSV、純文字）。"
    if any(d.name == name and d.size == len(data) for d in existing):
        return None, f"「{name}」已經在文件面板裡了。"
    if len(existing) >= MAX_DOCS:
        return None, f"一段對話最多放 {MAX_DOCS} 份文件，請先移除用不到的。"
    if sum(d.size for d in existing) + len(data) > MAX_TOTAL_BYTES:
        return None, f"文件合計不能超過 {MAX_TOTAL_BYTES // 1024 // 1024}MB，「{name}」放不進去。"
    note = ""
    try:
        if ext == ".pdf":
            doc, note = _pdf(name, data)
            if doc is None:
                return None, note
            used = sum(d.pages for d in existing)
            if used + doc.pages > MAX_PDF_PAGES:
                return None, (f"PDF 合計最多 {MAX_PDF_PAGES} 頁（目前已有 {used} 頁），"
                              f"「{name}」有 {doc.pages} 頁放不進去。可以先拆成小份再上傳。")
            return doc, note
        if ext in TEXT_EXTS:
            text = _decode(data).strip()
            if not text:
                return None, f"「{name}」裡找不到文字。"
            if len(text) > MAX_TEXT_CHARS:
                text, note = text[:MAX_TEXT_CHARS], f"「{name}」太長，只放得進前 {MAX_TEXT_CHARS:,} 字。"
            return Doc(name, "text", len(data), data=text), note
        if ext == ".docx":
            blocks, labels, n_paras = _docx(data)
            amount = f"{n_paras} 段" + (f"、{len(blocks) - n_paras} 個表格" if len(blocks) > n_paras else "")
        elif ext == ".pptx":
            blocks, labels = _pptx(data)
            amount = f"{len(blocks)} 張投影片"
        else:
            blocks, labels = _xlsx(data) if ext == ".xlsx" else _csv(data)
            amount = f"{sum(len(b.splitlines()) - 2 for b in blocks):,} 列資料"   # 每塊扣掉標籤和表頭
    except Exception as e:  # noqa: BLE001 — 壞檔或加密的 Office 檔
        return None, f"「{name}」打不開（{type(e).__name__}），請確認檔案沒有損壞或加密。"
    if not blocks:
        return None, f"「{name}」裡找不到文字。"
    blocks, labels, cut = _cap_blocks(blocks, labels)
    if cut:
        note = f"「{name}」太長，只放得進前面的 {labels[-1]}。"
    return Doc(name, "blocks", len(data), blocks=blocks, labels=labels, amount=amount), note


# ── 送進對話 ──

def _document_block(doc: Doc) -> dict:
    if doc.kind == "pdf":
        source: dict[str, Any] = {"type": "base64", "media_type": "application/pdf", "data": doc.data}
    elif doc.kind == "text":
        source = {"type": "text", "media_type": "text/plain", "data": doc.data}
    else:
        source = {"type": "content", "content": [{"type": "text", "text": b} for b in doc.blocks]}
    return {"type": "document", "title": doc.name, "source": source, "citations": {"enabled": True}}


def shelf_message(docs: list[Doc]) -> HumanMessage | None:
    """文件架：放在對話最前面。最後一份文件標 1h 快取斷點——它前面的系統提示、工具、所有文件一起快取。"""
    if not docs:
        return None
    blocks = [_document_block(d) for d in docs]
    blocks[-1] = {**blocks[-1], "cache_control": CACHE_DOCS}
    listing = "、".join(d.summary for d in docs)
    # 寫在文件旁邊比寫在系統提示有效：2026-10-24 本機 3 次有 2 次把英文原句整句貼進正文
    blocks.append({"type": "text", "text": f"（以上是使用者在文件面板上傳的 {len(docs)} 份文件：{listing}。"
                                           "之後的提問如果跟文件有關，請根據文件回答。正文一律用中文轉述，"
                                           "不要貼英文句子——英文原句系統會自動列在引用清單裡。）"})
    return HumanMessage(blocks, additional_kwargs={SHELF_KEY: True})


# 每題結尾的提醒（頁面併進 context_block，兩種模式都會在對話最後看到）。2026-10-24 線上與本機：第一題有引用，
# 追問時一條都沒有——歷史裡的前一則回答是純文字（看不到引用），模型照著自己上一則的樣子寫
TURN_NOTE = ("對話裡有使用者上傳的文件。這題如果根據文件回答，每個重點都要用引用標出出處"
             "（歷史裡的前一則回答看不到引用，不代表這次不用）；正文用中文轉述，不要貼英文原句。")


def is_shelf(msg: BaseMessage) -> bool:
    """有標記，或帶著快取斷點的文件區塊（中途被複製、標記沒跟著過去也認得出來）。"""
    if not isinstance(msg, HumanMessage):
        return False
    if msg.additional_kwargs.get(SHELF_KEY):
        return True
    return isinstance(msg.content, list) and any(
        isinstance(b, dict) and b.get("type") == "document" and b.get("cache_control") for b in msg.content)


def has_shelf(messages: list) -> bool:
    return any(is_shelf(m) for m in messages or [])


def signature(docs: list[Doc]) -> list[list]:
    """文件架的身分（順序有意義：引用的 document_index 照這個順序）。存進歷史用，所以是 list 不是 tuple。"""
    return [[d.name, d.size] for d in docs or []]


def cite_blocks(msg: AIMessage | None) -> list[dict]:
    """回答的文字區塊連同引用，存進歷史、下一題原樣傳回去。2026-10-24：前一則回答以純文字放進歷史時，
    追問常常一條引用都沒有（模型照著自己上一則的樣子寫）；保留引用結構後本機 6 次追問都有引用。"""
    if msg is None or isinstance(msg.content, str):
        return []
    return [{k: v for k, v in b.items() if k in ("type", "text", "citations")}
            for b in msg.content if isinstance(b, dict) and b.get("type") == "text" and b.get("text")]


# ── 引用 → 註腳 ──
DOC_CITATION_TYPES = ("page_location", "char_location", "content_block_location")
SNIPPET_CHARS = 80


def _location(c: dict, doc: Doc | None) -> str:
    t = c.get("type")
    if t == "page_location":
        s, e = int(c.get("start_page_number") or 1), int(c.get("end_page_number") or 0)
        return f"第 {s} 頁" if e - s <= 1 else f"第 {s}–{e - 1} 頁"        # end 不含
    if t == "content_block_location" and doc is not None and doc.labels:
        s, e = int(c.get("start_block_index") or 0), int(c.get("end_block_index") or 0)
        e = max(e - 1, s)                                                  # end 不含
        if s >= len(doc.labels):
            return ""
        return doc.labels[s] if e == s or e >= len(doc.labels) else f"{doc.labels[s]}～{doc.labels[e]}"
    return ""


class Citer:
    """把回答裡的文件引用編號：被引用的段落後面加上 [n]，同一段引文重複引用沿用同一號。"""

    def __init__(self, docs: list[Doc]):
        self.docs = list(docs)
        self.notes: list[str] = []
        self._keys: dict[tuple, int] = {}

    def __call__(self, block: dict) -> str:
        nums = []
        for c in block.get("citations") or []:
            if not isinstance(c, dict) or c.get("type") not in DOC_CITATION_TYPES:
                continue
            i = c.get("document_index")
            doc = self.docs[i] if isinstance(i, int) and 0 <= i < len(self.docs) else None
            name = doc.name if doc else (c.get("document_title") or "文件")
            loc = _location(c, doc)
            # PDF 斷字會留下看不見的控制字元（2026-10-24 線上：「Spi\x02derMonkey」）
            quote = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", c.get("cited_text") or "")
            quote = re.sub(r"^【[^】]*】", "", re.sub(r"\s+", " ", quote).strip()).strip()   # 區塊開頭的位置標籤
            key = (name, loc, quote)
            if key not in self._keys:
                self._keys[key] = len(self._keys) + 1
                snippet = quote[:SNIPPET_CHARS] + ("…" if len(quote) > SNIPPET_CHARS else "")
                self.notes.append(f"{self._keys[key]}. {name}{(' ' + loc) if loc else ''}："
                                  f"「{snippet.replace('[', '［').replace(']', '］')}」")
            if self._keys[key] not in nums:
                nums.append(self._keys[key])
        return "".join(f"<sup>[{n}]</sup>" for n in nums)

    def footer(self) -> str:
        """引用清單（系統產生，不靠模型寫）；history.strip_sources_footer 會在送回模型前拿掉。"""
        if not self.notes:
            return ""
        return "\n\n---\n**📄 文件引用**\n" + "\n".join(self.notes)


def cited_answer(msg: AIMessage | None, docs: list[Doc]) -> tuple[str, str]:
    """(加了 [n] 的回答文字, 引用清單)。沒有文件時就是原本的回答文字。"""
    from .events import answer_text
    if not docs:
        return answer_text(msg), ""
    citer = Citer(docs)
    return answer_text(msg, mark=citer), citer.footer()
